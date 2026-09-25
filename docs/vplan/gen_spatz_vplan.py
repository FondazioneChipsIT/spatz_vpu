#!/usr/bin/env python3
# Generates SPATZ_VERIFICATION_PLAN.xlsx (sheets "Verification Plan" and "Configuration").
#
# Every Note is a verbatim quote of a source file; the script aborts if a quote
# is not found in its source, so the Note column cannot drift from the spec.
#
# Usage (from working_dir/spatz_vpu): python3 docs/vplan/gen_spatz_vplan.py
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]          # working_dir/spatz_vpu
CLUSTER = ROOT.parents[1]                             # spatz repo root
OUT = ROOT / "docs" / "vplan" / "SPATZ_VERIFICATION_PLAN.xlsx"

SOURCES = {
    "SPATZ_SPEC.md": ROOT / "docs" / "SPATZ_SPEC.md",
    "README.md": ROOT / "README.md",
    "spatz_controller.sv": ROOT / "hw" / "src" / "spatz_controller.sv",
    "spatz_vlsu.sv": ROOT / "hw" / "src" / "spatz_vlsu.sv",
    "spatz_pkg.sv.tpl": ROOT / "hw" / "src" / "spatz_pkg.sv.tpl",
    "spatz.sv": ROOT / "hw" / "src" / "spatz.sv",
}
_cache = {}


def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def quote(src, text):
    """Return a Note cell: verbatim quote + source. Abort if not verbatim."""
    if src not in _cache:
        _cache[src] = norm(SOURCES[src].read_text())
    if norm(text) not in _cache[src]:
        sys.exit(f"[vplan] quote not found verbatim in {src}:\n  {text}")
    return f"{src}: \"{text}\""


def notes(*qs):
    return "\n".join(quote(s, t) for s, t in qs)


# ----------------------------------------------------------------------------
# Verification Plan sheet
# ----------------------------------------------------------------------------
# Old testbench = cluster-level SW regression: self-checking C tests
# (sw/riscvTests, sw/spatzBenchmarks) run on the spatz_cluster testbench by
# util/local_ci.sh; results are compared in software with golden values after
# a vector store (VCMP/VMCMP macros), pass/fail is the program exit code.
OLD_GEN = ("Cluster-level SW regression (util/local_ci.sh on spatz_cluster TB, "
           "self-checking C tests: golden values compared in SW after vse*). ")

S = "SPATZ_SPEC.md"
C_ALL = ""                   # empty = all legal configurations
C_DEF = "CS-DEF"
C_CI = "CS-DEF, CS-SVRF, CS-32B, CS-DBW"

items = [
    # --- Interfaces --------------------------------------------------------
    dict(
        f="Offload issue handshake",
        goal="Verify that every instruction presented on issue_* is consumed exactly once when "
             "issue_valid_i && issue_ready_o, and that the DUT never loses or duplicates an "
             "instruction while issue_ready_o is low.",
        test="-",
        stim="Issue agent drives a constrained-random legal instruction stream with random idle "
             "cycles between instructions, keeps issue_valid_i/issue_req_i stable while "
             "issue_ready_o=0; mix instructions executed locally in the FPU sequencer (fmv.*, "
             "flw/fsw) with forwarded ones, and FP instructions whose FP sources are still "
             "pending (sequencer stall).",
        status="Not Started",
        note=notes((S, "The core asserts `issue_valid_i` with a stable `issue_req_i` until "
                       "`issue_ready_o` is seen."),
                   (S, "An illegal instruction **is consumed** (`issue_ready_o` = 1) and "
                       "answered with `accept = 0, exception = 1`; it is not executed.")),
        old=OLD_GEN + "Covered: the Snitch core drives the interface with real programs, so "
            "the nominal handshake is exercised. Not covered: no protocol checker, no random "
            "valid gaps, no control over when the core presents instructions.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="Issue immediate response",
        goal="Verify that in the handshake cycle issue_rsp_o reports accept/writeback/loadstore/"
             "exception consistently with the class of the accepted instruction.",
        test="-",
        stim="For every instruction class of the Supported Instructions tables (CON, VFU, LSU, "
             "SLD, SEQ, illegal) issue at least one instruction; sample issue_rsp_o only when "
             "issue_valid_i && issue_ready_o.",
        status="Not Started",
        note=notes((S, "`issue_rsp_o` (except `isfloat`) is only meaningful in the cycle where "
                       "`issue_valid_i && issue_ready_o`."),
                   (S, "The Snitch core only uses `loadstore` and `isfloat`; `accept`, "
                       "`writeback` and `exception` are exposed for completeness and for "
                       "checking.")),
        old=OLD_GEN + "Not covered: Snitch ignores accept/writeback/exception, no check exists; "
            "loadstore is only checked indirectly (a wrong value breaks scalar/vector memory "
            "ordering).",
        cfg=C_ALL, check="scoreboard"),
    dict(
        f="isfloat busy indication",
        goal="Verify that issue_rsp_o.isfloat is high whenever at least one vector instruction "
             "is in flight and low otherwise.",
        test="-",
        stim="Issue single vector instructions and bursts up to NrParallelInstructions, with and "
             "without memory backpressure; issue CSR-only sequences (isfloat must stay low).",
        status="Not Started",
        note=notes((S, "`issue_rsp_o.isfloat` is a level signal, high whenever at least one "
                       "vector instruction is in flight.")),
        old=OLD_GEN + "Not covered: used by Snitch only to delay fcsr reads; no test reads fcsr "
            "while vector instructions run.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="Write-back response channel",
        goal="Verify that exactly one rsp_o transfer is produced per write-back instruction, in "
             "issue order, with rsp_o.id equal to the request id and the correct data, and that "
             "rsp_valid_o/rsp_o stay stable under rsp_ready_i backpressure.",
        test="riscvTests-vsetvli, riscvTests-scalar",
        stim="Issue vsetvl*/CSR reads (including rd=x0), vmv.x.s, scalar mul/div/rem, FP compares/"
             "fclass/fcvt.w* and fmv.x.* back to back; drive rsp_ready_i with random "
             "backpressure, including long stalls.",
        status="In Progress",
        note=notes((S, "every `vsetvl`/`vsetvli`/`vsetivli` (data = new `vl`) and every vector "
                       "CSR access (data = CSR read value), **regardless of `rd`** (also when "
                       "`rd = x0`);"),
                   (S, "`rsp_o` is driven by a non-bypassable spill register: `rsp_valid_o` and "
                       "`rsp_o` are registered outputs and remain stable until `rsp_ready_i`.")),
        old=OLD_GEN + "Covered: returned values used by the C code (vl, scalar results) are "
            "checked implicitly. Not covered: ordering and id checks, rd=x0 responses, "
            "backpressure (the core accepts responses without stalls).",
        cfg=C_ALL, check="scoreboard + assertion"),
    dict(
        f="FP-destination results kept internal",
        goal="Verify that results with an FP destination are written to the internal FP "
             "register file and never appear on rsp_o (rsp_o.id[5] always 0).",
        test="riscvTests-scalar",
        stim="Issue vfmv.f.s, scalar FP arithmetic, flw/fld and fmv.w.x, then read the FP "
             "registers back with fmv.x.w / fsw.",
        status="In Progress",
        note=notes((S, "Results with an FP destination (`vfmv.f.s`, scalar FP arithmetic, FP "
                       "loads, `fmv.{b,h,w}.x`) are written to the internal FP register file "
                       "and **do not appear on `rsp_o`**.")),
        old=OLD_GEN + "Covered: FP scalar values are used by later instructions whose results "
            "are checked. Not covered: no check that nothing leaks on rsp_o.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="VLSU memory port protocol",
        goal="Verify that each VLSU port issues ELEN-aligned requests with amo=AMONone, keeps "
             "valid/payload stable until ready, and operates correctly with in-order responses "
             "of any latency, including one response per store.",
        test="-",
        stim="Memory slave agents on every port: random req_ready deassertion (single cycles and "
             "long stalls), random in-order response latency 1..32 cycles, response for every "
             "write; run all load/store addressing modes on top.",
        status="In Progress",
        note=notes((S, "**Responses must be returned in request order per port** (non-MemPool "
                       "build: the \"reorder buffer\" is a FIFO)."),
                   (S, "**Every request, including stores, must receive exactly one response** "
                       "(`spatz_mem_rsp_valid_i` pulse).")),
        old=OLD_GEN + "Covered: real TCDM traffic, data correctness checked end to end. Not "
            "covered: TCDM has fixed latency, backpressure only from natural bank conflicts, "
            "no protocol assertions.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="VLSU outstanding-request limits [assumed]",
        goal="Verify that no port has more than 8 loads outstanding/buffered and that the DUT "
             "behaves correctly with up to 7 unacknowledged stores per port.",
        test="-",
        stim="Directed: long unit-stride and strided stores/loads with memory latency > 8 "
             "cycles and no ready backpressure; back-to-back store->load on the same port.",
        status="Not Started",
        note=notes((S, "For loads the VLSU never has more than 8 requests per port outstanding "
                       "or buffered (bounded by its offset queue), so its buffers cannot "
                       "overflow."),
                   (S, "Stores are **not** limited by the VLSU; the per-port store-acknowledge "
                       "counter is only 3 bits wide, so the memory must not keep more than 7 "
                       "stores per port unacknowledged.")),
        old=OLD_GEN + "Not covered: TCDM latency is 1-2 cycles, the limits are never reached.",
        cfg=C_ALL, check="assertion + directed test"),
    dict(
        f="FP LSU (reqrsp) interface",
        goal="Verify that scalar FP loads/stores produce correct reqrsp requests (address, size, "
             "strb, NaN-boxing of loaded values) and complete with in-order responses under "
             "backpressure on q and p channels.",
        test="riscvTests-fl_narrow, riscvTests-scalar",
        stim="flb/flh/flw/fld and fsb/fsh/fsw/fsd with all alignments; random q_ready and "
             "p_valid delays; up to NumOutstandingLoads loads in flight.",
        status="In Progress",
        note=notes((S, "Standard `reqrsp` protocol: two independent valid/ready channels (`q` "
                       "and `p`), responses in order, one response per request (including "
                       "stores).")),
        old=OLD_GEN + "Covered: FP scalar loads in several tests. Not covered: backpressure, "
            "outstanding limit, protocol assertions.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="Memory completion side channel",
        goal="Verify that spatz_mem_finished_o[1]/[0] pulse exactly once per completed "
             "vector/FP memory instruction (last segment for segmented ones) and "
             "spatz_mem_str_finished_o only together with finished for stores.",
        test="-",
        stim="Sequences of vector loads, stores and segmented accesses interleaved with FP "
             "loads/stores, with memory backpressure; count pulses per instruction.",
        status="Not Started",
        note=notes((S, "`spatz_mem_finished_o[1]` pulses once per completed vector memory "
                       "instruction (for segmented instructions only on the last segment); "
                       "`spatz_mem_str_finished_o[1]` pulses together with it when that "
                       "instruction was a store.")),
        old=OLD_GEN + "Covered only indirectly: Snitch uses the pulses to order scalar memory "
            "accesses. Not covered: no direct count check.",
        cfg=C_ALL, check="assertion + scoreboard"),
    dict(
        f="Scalar FP vs vector memory ordering",
        goal="Verify that FP loads/stores executed by the FPU sequencer observe the memory "
             "effects of preceding vector stores and are not overtaken by later vector "
             "accesses.",
        test="-",
        stim="Directed: vse32 then flw from the same address; vle32 then fsw to an address in "
             "the loaded range; repeat with long memory latency on the VLSU ports.",
        status="Not Started",
        note=notes((S, "FP loads/stores are delayed while vector memory instructions of the "
                       "conflicting type are in flight (tracked with `spatz_mem_finished`), "
                       "with a 3-bit counter of in-flight vector memory instructions.")),
        old=OLD_GEN + "Not covered: tests do not deliberately mix FP scalar and vector accesses "
            "to the same address.",
        cfg=C_ALL, check="scoreboard + directed test"),
    dict(
        f="FPU side channel sampling",
        goal="Verify that fpu_rnd_mode_i and fpu_fmt_mode_i are sampled when the instruction is "
             "accepted and that later changes do not affect instructions already in flight.",
        test="-",
        stim="Issue FP vector instructions and change fpu_rnd_mode_i / fpu_fmt_mode_i in the "
             "cycle after the handshake while they are still executing.",
        status="Not Started",
        note=notes((S, "`fpu_rnd_mode_i` and `fpu_fmt_mode_i` are \"untimed\": they are sampled "
                       "by the decoder in the cycle the instruction is accepted and travel with "
                       "the instruction.")),
        old=OLD_GEN + "Not covered: no test changes frm or the FP format mode.",
        cfg=C_ALL, check="scoreboard + directed test"),
    # --- Programming model -----------------------------------------------
    dict(
        f="Vector CSRs",
        goal="Verify the read value and write semantics of vstart, vxsat, vxrm, vcsr, vl, "
             "vtype, vlenb for csrrw/s/c[i], and that any other CSR address is illegal.",
        test="riscvTests-vsetvli",
        stim="Directed: every CSR x every csr* variant with rd=x0/rd!=x0 and rs1=x0/rs1!=x0; "
             "reads after vsetvl*; accesses to non-vector CSR addresses.",
        status="In Progress",
        note=notes((S, "Any other CSR address offloaded to Spatz is illegal."),
                   (S, "Reads as 0, writes ignored (no fixed-point support)")),
        old=OLD_GEN + "Covered: vl returned by vsetvli. Not covered: no test reads or writes "
            "the vector CSRs directly.",
        cfg=C_ALL, check="scoreboard + directed test"),
    dict(
        f="vsetvl* configuration",
        goal="Verify vl/vtype after vsetvli/vsetivli/vsetvl for all SEW x LMUL (incl. "
             "fractional), all AVL cases, rs1=x0 rules, keep-vl ratio check and every illegal "
             "vtype.",
        test="riscvTests-vsetvli, riscvTests-vill",
        stim="Sweep SEW {8,16,32,64} x LMUL {F8,F4,F2,1,2,4,8,reserved}, AVL {0,1,VLMAX-1,VLMAX,"
             "VLMAX+1,max}, rd/rs1 = x0 combinations, keep-vl with same and different "
             "SEW/LMUL ratio.",
        status="In Progress",
        note=notes((S, "If `rs1 = x0` and `rd = x0` (*keep-vl*): `vl` is preserved; if the new "
                       "`vtype` changes the `SEW/LMUL` ratio, `vtype` is set to illegal."),
                   (S, "`VLMAX = (VLENB >> vsew) << vlmul` (right-shift for fractional LMUL). "
                       "`vl = min(AVL, VLMAX)`.")),
        old=OLD_GEN + "Covered: vsetvli with a set of SEW/LMUL values, one illegal vtype (vill "
            "test, expected to fail). Not covered: fractional LMUL, keep-vl, full AVL sweep.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="vstart handling",
        goal="Verify that execution starts at vstart for arithmetic, memory and slide "
             "instructions and that vstart is reset to 0 by every accepted non-CSR instruction.",
        test="-",
        stim="csrw vstart with values {1, word-1, word, vl-1} before each instruction class; "
             "read vstart back after the instruction.",
        status="Not Started",
        note=notes((S, "Reset to 0 by every accepted non-CSR instruction (including "
                       "`vsetvl*`)")),
        old=OLD_GEN + "Not covered: no enabled test writes vstart.",
        cfg=C_ALL, check="scoreboard + directed test"),
    dict(
        f="Masking (vm = 0)",
        goal="Verify that inactive elements are left undisturbed for arithmetic, loads, stores "
             "and slides, and that comparisons write inactive mask bits according to vtype.vma.",
        test="riscvTests-* (98 of 115 enabled tests use v0.t)",
        stim="Every supported instruction with vm=0 and random v0 (all-zero, all-one, "
             "alternating, random), both vma values for comparisons.",
        status="In Progress",
        note=notes((S, "Inactive elements are **mask-undisturbed** for arithmetic, loads and "
                       "slides (not written: byte enables are cleared)."),
                   (S, "For comparisons (mask-producing), inactive bits follow `vtype.vma`: "
                       "`vma = 1` → written as 1 (agnostic), `vma = 0` → undisturbed (the old "
                       "`vd` is read first).")),
        old=OLD_GEN + "Covered: masked variants in most instruction tests with fixed masks. Not "
            "covered: random masks, vma=0 on comparisons, masked slides.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Tail policy [assumed]",
        goal="Verify that tail elements (index >= vl) are never modified, for both vta values.",
        test="-",
        stim="Pre-fill destination registers with a known pattern, run every instruction class "
             "with vl not a multiple of the VRF word and vta in {0,1}, store the full register "
             "group.",
        status="Not Started",
        note=notes((S, "Tail elements are **always undisturbed** (byte enables limited to "
                       "`vl`), independent of `vtype.vta`.")),
        old=OLD_GEN + "Not covered: tests compare only the first vl elements.",
        cfg=C_ALL, check="scoreboard"),
    # --- Instructions ------------------------------------------------------
    dict(
        f="Integer vector arithmetic",
        goal="Verify the result of every supported integer instruction (add/sub/logic/shift/"
             "min/max/merge) for all SEW, LMUL and vl classes.",
        test="riscvTests-vadd ... riscvTests-vmaxu (see Old testbench situation)",
        stim="Constrained-random operands (including min/max values), .vv/.vx/.vi forms, SEW "
             "x LMUL x vl {1, <word, =word, not multiple of word, VLMAX}.",
        status="In Progress",
        note=notes((S, "`vadd.{vv,vx,vi}`, `vsub.{vv,vx}`, `vrsub.{vx,vi}`")),
        old=OLD_GEN + "Covered: vadd, vsub, vrsub, vand, vor, vxor, vsll, vsrl, vsra, vmin[u], "
            "vmax[u], vmv with directed operands, LMUL 1..8. Not covered: fractional LMUL, "
            "random operands, vmerge (no enabled test).",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Widening integer arithmetic",
        goal="Verify vwadd[u], vwsub[u], vwmul[u,su] and vwmacc[u,su,us] results and "
             "destination register group placement.",
        test="riscvTests-vwadd, -vwaddu, -vwsub, -vwsubu, -vwmul, -vwmulu, -vwmulsu, -vwmacc, "
             "-vwmaccu, -vwmaccsu, -vwmaccus",
        stim="Signed/unsigned extremes, all SEW < ELEN, LMUL up to 4, masked and unmasked.",
        status="In Progress",
        note=notes((S, "`vwadd[u].{vv,vx}`, `vwsub[u].{vv,vx}` (no `.wv/.wx` forms)")),
        old=OLD_GEN + "Covered: one test per instruction with directed values. Not covered: "
            "random operands, vl not multiple of the word.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Integer multiply, multiply-add and divide",
        goal="Verify vmul*, vmacc/vnmsac/vmadd/vnmsub and vdiv[u]/vrem[u] including division "
             "by zero and signed overflow per RVV, independently of the serial divider latency.",
        test="riscvTests-vmul, -vmulh, -vmulhu, -vmulhsu, -vmacc, -vmadd, -vnmsac, -vnmsub, "
             "-vdiv, -vdivu, -vrem, -vremu",
        stim="Random operands plus directed divisor=0, INT_MIN/-1, operands with very different "
             "leading-zero counts (divider latency extremes).",
        status="In Progress",
        note=notes((S, "`vdiv[u]/vrem[u].{vv,vx}` (serial divider, multi-cycle)")),
        old=OLD_GEN + "Covered: directed tests per instruction. Not covered: RVV corner cases "
            "not confirmed in the tests, random operands.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Add/subtract with carry [open question Q03]",
        goal="Verify vadc/vsbc/vmadc/vmsbc results with carry/borrow taken from v0.",
        test="-",
        stim="Random operands and random v0 for vvm/vxm/vim forms; carry-out forms with and "
             "without carry-in.",
        status="Not Started",
        note=notes((S, "`vadc.{vvm,vxm,vim}`, `vmadc.{vv,vx,vi,vvm,vxm,vim}`, `vsbc.{vvm,vxm}`, "
                       "`vmsbc.{vv,vx,vvm,vxm}`")),
        old=OLD_GEN + "Not covered: vmadc.c / vsbc.c exist in sw/riscvTests/isa/rv64uv but are "
            "not enabled in CMakeLists.txt.",
        cfg=C_ALL, check="scoreboard"),
    dict(
        f="Integer reductions",
        goal="Verify vred* results (element 0 of vd) for all SEW, vl and masks, including the "
             "neutral value for masked-off elements.",
        test="riscvTests-vredsum, -vredand, -vredor, -vredxor, -vredmin, -vredminu, -vredmax, "
             "-vredmaxu",
        stim="vl {1, 2, word, not multiple of word, VLMAX}, masked/unmasked, vs1[0] extremes.",
        status="In Progress",
        note=notes((S, "Reductions are executed by a dedicated FSM: initialisation with "
                       "`vs1[0]`, accumulation of the whole vector per lane, intra-lane drain, "
                       "inter-lane tree, SIMD fold, write-back of element 0 of `vd`.")),
        old=OLD_GEN + "Covered: one test per reduction. Not covered: vl sweep, random masks.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Compares and mask-logical instructions",
        goal="Verify integer/FP compares and vm*.mm results, the number of mask bytes written "
             "(ceil(vl/8)) and the vma policy.",
        test="riscvTests-vmfeq, -vmfne, -vmflt, -vmfle, -vmfgt, -vmfge, -vmand, -vmor, "
             "-vmandnot, -vmnand, -vmnor, -vmornot, -vmxnor, -vmxor",
        stim="All compare forms with equal/greater/less/NaN operands, vl not multiple of 8, "
             "vma in {0,1}.",
        status="In Progress",
        note=notes((S, "For mask-producing instructions, the whole mask bytes up to "
                       "`ceil(vl/8)` are written.")),
        old=OLD_GEN + "Covered: FP compares and mask-logical ops. Not covered: integer compares "
            "(vmseq/vmsne/vmslt[u]/vmsle[u]/vmsgt[u] have no enabled test), vma=0.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Unit-stride loads and stores",
        goal="Verify vle*/vse* for all EEW, aligned and unaligned base addresses, masked, with "
             "the port interleaving and strobes defined in the spec.",
        test="riscvTests-vle8, -vle16, -vle32, -vle64, -vse8, -vse16, -vse32, -vse64",
        stim="Base address offset 0..7, vl classes, LMUL 1..8, masked; memory agents with random "
             "backpressure; check per-port address sequence.",
        status="In Progress",
        note=notes((S, "**Unit-stride, aligned base, `vstart = 0`:** each port transfers a full "
                       "8-byte word per request.")),
        old=OLD_GEN + "Covered: one test per EEW. Not covered: unaligned base sweep, port "
            "address sequence, backpressure.",
        cfg=C_CI, check="scoreboard + assertion"),
    dict(
        f="Strided loads and stores",
        goal="Verify vlse*/vsse* for positive, negative, zero and non-element-multiple strides.",
        test="riscvTests-vls, riscvTests-vss",
        stim="Stride in {0, SEW/8, -SEW/8, large, not multiple of SEW/8}, all EEW, masked.",
        status="In Progress",
        note=notes((S, "Address = `rs1 + offset`, offset = `i * stride` (strided, `stride = "
                       "rs2`)")),
        old=OLD_GEN + "Covered: vls/vss with fixed strides. Not covered: negative, zero and "
            "non-multiple strides (open question Q13).",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Indexed loads and stores",
        goal="Verify vl[ou]xei*/vs[ou]xei* for every index EEW x data SEW combination, with "
             "sign-extended byte offsets.",
        test="riscvTests-vloxei, riscvTests-vsuxei",
        stim="Index EEW {8,16,32,64} x SEW {8,16,32,64}, random and repeated indices, "
             "negative offsets, masked.",
        status="In Progress",
        note=notes((S, "ordered and unordered handled identically (sequential). Index EEW 64 "
                       "uses only the low 32 bits")),
        old=OLD_GEN + "Covered: vloxei and vsuxei tests. Not covered: full EEW x SEW cross, "
            "vluxei/vsoxei encodings, negative offsets.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Segmented loads and stores",
        goal="Verify vlseg<nf>/vsseg<nf> for nf 2..8 and all EEW, writing register vd+field "
             "for each field.",
        test="riscvTests-vlseg, riscvTests-vsseg",
        stim="nf 2..8 x EEW x LMUL with nf*LMUL <= 8, masked, unaligned base.",
        status="In Progress",
        note=notes((S, "`nf` = 2..8; register `vd + field` per field")),
        old=OLD_GEN + "Covered: vlseg/vsseg tests. Not covered: full nf x EEW sweep.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="Permutation (VSLDU)",
        goal="Verify slides (any amount), vslide1*/vfslide1*, vmv.* and vcompress.vm results, "
             "masked and unmasked.",
        test="riscvTests-vslideup, -vslidedown, -vslide1up, -vslide1down, -vfslide1up, "
             "-vfslide1down, -vmv, -vfmv",
        stim="Slide amount {0, 1, word-1, word, >vl, >VLMAX}, all SEW, LMUL up to 8; vcompress "
             "with random masks.",
        status="In Progress",
        note=notes((S, "`vmv.v.{v,x,i}`, `vmv.s.x`, `vslideup.{vx,vi}`, `vslide1up.vx`, "
                       "`vslidedown.{vx,vi}`, `vslide1down.vx`, `vfslide1up.vf`, "
                       "`vfslide1down.vf`, `vfmv.v.f`, `vfmv.s.f`, `vcompress.vm`.")),
        old=OLD_GEN + "Covered: slides and moves. Not covered: vcompress (vcompress.c not "
            "enabled), extreme slide amounts.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="FP vector arithmetic",
        goal="Verify FP add/sub/mul/FMA/min/max/sgnj results for every format (incl. FP16ALT/"
             "FP8ALT via fpu_fmt_mode_i), every rounding mode and IEEE special operands.",
        test="riscvTests-vfadd, -vfsub, -vfrsub, -vfmul, -vfmacc, -vfnmacc, -vfmsac, -vfnmsac, "
             "-vfmadd, -vfnmadd, -vfmsub, -vfnmsub, -vfmin, -vfmax, -vfsgnj, -vfsgnjn, -vfsgnjx",
        stim="Random operands plus ±0, ±inf, qNaN, sNaN, subnormals, rounding ties; "
             "fpu_rnd_mode_i in {RNE,RTZ,RDN,RUP,RMM}; fpu_fmt_mode_i src/dst in {0,1}.",
        status="In Progress",
        note=notes((S, "Supported FP formats (default, `RVD = 1`): FP64, FP32, FP16, FP16ALT "
                       "(bfloat16), FP8, FP8ALT, selected by SEW and `fpu_fmt_mode_i`")),
        old=OLD_GEN + "Covered: directed values in FP64/FP32/FP16 with the default rounding "
            "mode. Not covered: rounding modes (no test changes frm), ALT formats, special "
            "operands sweep.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="FP conversions",
        goal="Verify vfcvt.* and vfncvt.* results including saturation, NaN handling and RTZ "
             "variants.",
        test="riscvTests-vfcvt, riscvTests-vfncvt",
        stim="Values out of integer range, NaN, ±inf, subnormals; every SEW pair; all rounding "
             "modes.",
        status="In Progress",
        note=notes((S, "`vfcvt.{f.x,f.xu,x.f,xu.f,rtz.x.f,rtz.xu.f}.v`, `vfncvt.{f.x,f.xu,x.f,"
                       "xu.f,rtz.x.f,rtz.xu.f,f.f}.w`")),
        old=OLD_GEN + "Covered: vfcvt/vfncvt directed tests. Not covered: saturation and "
            "rounding-mode sweep.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="FP widening and dot product [open question Q05]",
        goal="Verify vfw* results and vfwdotp (SDOTP) accumulation into vd, including special "
             "operands of the narrow format.",
        test="riscvTests-vfwadd, -vfwsub, -vfwmul, -vfwmacc, -vfwmsac, -vfwnmsac; "
             "spatzBenchmarks sdotp-*",
        stim="All narrow formats, .vv/.vf/.wv/.wf forms, subnormal/inf/NaN narrow operands.",
        status="In Progress",
        note=notes((S, "Widening FP instructions widen the operands before the FPU with a simple "
                       "exponent re-bias (`widen_fp*_to_fp*` in `spatz_pkg`), not through an "
                       "FPnew conversion")),
        old=OLD_GEN + "Covered: vfw* directed tests, vfwdotp through benchmarks only. Not "
            "covered: special narrow operands, vfwnmacc (no enabled test).",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="FP reductions [open question Q06]",
        goal="Verify vfredusum/vfredosum/vfredmin/vfredmax results and the neutral value used "
             "for masked-off elements.",
        test="riscvTests-vfredusum, -vfredosum, -vfredmin, -vfredmax",
        stim="vl sweep, masked, operands chosen so that summation order changes the result.",
        status="In Progress",
        note=notes((S, "`vfredosum/vfredusum/vfredmin/vfredmax.vs` (ordered sum is executed as a "
                       "tree, like the unordered one)")),
        old=OLD_GEN + "Covered: one test per reduction. Not covered: order-sensitive operands.",
        cfg=C_CI, check="scoreboard"),
    dict(
        f="FP exception flags",
        goal="Verify that the flags OR-accumulated from fpu_status_o over a sequence equal the "
             "reference fflags.",
        test="-",
        stim="Sequences producing each flag (NV, DZ, OF, UF, NX) in isolation and combined, "
             "vector and scalar, with and without masked/tail lanes.",
        status="Not Started",
        note=notes((S, "`fpu_status_o` is the OR of the registered FPnew `status_o` of all FPUs, "
                       "updated every cycle; it is intended to be OR-accumulated (sticky) into "
                       "`fflags`, not to be sampled per instruction.")),
        old=OLD_GEN + "Not covered: no test reads fflags.",
        cfg=C_ALL, check="scoreboard"),
    dict(
        f="Scalar instructions executed by Spatz",
        goal="Verify offloaded mul/div/rem, scalar FP arithmetic/compare/convert, fmv.* and "
             "FP loads/stores.",
        test="riscvTests-scalar, riscvTests-fl_narrow",
        stim="Random operands for each scalar instruction, interleaved with vector "
             "instructions to exercise the IPU0/FPU0 sharing.",
        status="In Progress",
        note=notes((S, "Snitch offloads its M-extension and its whole F/D/Zfh/Xf8 extension to "
                       "Spatz.")),
        old=OLD_GEN + "Covered: scalar and fl_narrow tests. Not covered: interleaving with "
            "vector instructions.",
        cfg=C_CI, check="scoreboard"),
    # --- Errors ------------------------------------------------------------
    dict(
        f="Illegal instructions [open questions Q01, Q02]",
        goal="Verify that every unsupported encoding/CSR, unsupported EEW, FP div/sqrt and "
             "vector instruction with vill=1 is answered with accept=0/exception=1 and has no "
             "architectural side effect (no VRF, memory or rsp_o change).",
        test="riscvTests-invalid, riscvTests-vill",
        stim="Directed list of illegal encodings per class; illegal instruction issued while "
             "rsp_ready_i=0 and the response register is full; vector instructions after an "
             "illegal vsetvl.",
        status="In Progress",
        note=notes((S, "Reasons for `exception = 1`: unknown encoding, unsupported instruction "
                       "or CSR, invalid memory element width, FP instruction without FPU, `vill "
                       "= 1` for VFU vector / LSU / SLD instructions, FP division/sqrt.")),
        old=OLD_GEN + "Covered: invalid and vill tests, marked WILL_FAIL in CMakeLists.txt. Not "
            "covered: the failure mechanism is not asserted (open question Q14); no check of "
            "missing side effects.",
        cfg=C_ALL, check="scoreboard + directed test"),
    # --- Micro-architecture ------------------------------------------------
    dict(
        f="Dependencies and chaining",
        goal="Verify that results equal sequential execution for RAW/WAR/WAW hazards across "
             "VFU, VLSU and VSLDU, with up to NrParallelInstructions in flight and with "
             "chaining-inhibited producers.",
        test="riscvTests-vls_chain; spatzBenchmarks",
        stim="Random instruction streams with forced register reuse (same vd/vs), producer/"
             "consumer pairs for each unit pair, vslideup/strided/indexed/segmented producers.",
        status="In Progress",
        note=notes((S, "Every VRF port request carries the instruction ID; the port is enabled "
                       "only if all dependencies of that instruction wrote the VRF in the "
                       "previous cycle."),
                   (S, "Chaining is disabled for \"risky\" producers: `vslideup`, strided, "
                       "segmented and indexed memory instructions.")),
        old=OLD_GEN + "Covered: vls_chain and benchmark kernels exercise chaining incidentally. "
            "Not covered: hazards are not generated by construction.",
        cfg=C_CI, check="scoreboard + constrained-random regression"),
    dict(
        f="VRF bank arbitration and BUF_FPU",
        goal="Verify that VRF bank conflicts between units and a full BUF_FPU FIFO only delay "
             "execution and never corrupt results.",
        test="spatzBenchmarks (dp-fmatmul, sp-fmatmul, ...)",
        stim="Concurrent VFU and VLSU writes to registers mapped to the same bank; long FMA "
             "sequences with memory backpressure to fill the FPU buffer.",
        status="In Progress",
        note=notes((S, "Write arbitration per bank, fixed priority: `VFU > VLSU > VSLDU` by "
                       "default.")),
        old=OLD_GEN + "Covered: benchmarks create conflicts incidentally. Not covered: "
            "conflicts by construction, BUF_FPU disabled (buf_fpu=1 in every configuration).",
        cfg="CS-DEF, CS-NOBUF", check="scoreboard"),
    dict(
        f="Reset",
        goal="Verify CSR reset values (vtype.vill=1, vl=0, vstart=0), isfloat=0 and no memory "
             "request after reset, including reset asserted during activity.",
        test="-",
        stim="Directed: read CSRs after reset; assert rst_ni asynchronously in the middle of "
             "vector load/store and FPU sequences, then run a clean sequence.",
        status="Not Started",
        note=notes((S, "Reset values: `vstart = 0`, `vl = 0`, `vtype = {vill: 1, vsew: EW_8, "
                       "vlmul: LMUL_1, others 0}`."),
                   (S, "The VRF is **not reset** (latch array): its content after reset is "
                       "undefined (X in simulation).")),
        old=OLD_GEN + "Not covered: only power-on reset, never reset during activity.",
        cfg=C_ALL, check="assertion + directed test"),
    dict(
        f="Liveness under backpressure",
        goal="Verify that every accepted instruction eventually retires under any legal "
             "backpressure on rsp_ready_i, memory ready and FP LSU channels.",
        test="-",
        stim="Constrained-random streams with long random stalls on every slave interface; "
             "bounded stall windows followed by release.",
        status="Not Started",
        note=notes((S, "There is no response `ready`: the memory model must never expect "
                       "backpressure on responses.")),
        old=OLD_GEN + "Not covered: no controllable backpressure.",
        cfg=C_ALL, check="assertion + constrained-random regression"),
    # --- Parameters / configurations --------------------------------------
    dict(
        f="Parameter rules",
        goal="Verify that illegal parameter combinations are rejected at elaboration "
             "(NrMemPorts != N_FU, non-power-of-two N_FU/VLEN, VLEN < N_FU*ELEN, memory data "
             "width != ELEN, N_IPU = 0).",
        test="-",
        stim="Elaboration-only runs with each illegal combination; check the $error message.",
        status="Not Started",
        note=notes(("spatz.sv", "$error(\"[spatz] The number of FUs needs to be a power of "
                                "two\");"),
                   ("spatz_vlsu.sv", "$error(\"[spatz_vlsu] The number of memory ports needs "
                                     "to be equal to the number of FUs.\");")),
        old="No elaboration tests.",
        cfg="Illegal combinations only", check="directed test (elaboration) + review"),
    dict(
        f="Configuration space",
        goal="Verify that the functional items hold in every configuration set of the "
             "Configuration sheet.",
        test="util/local_ci.sh (riscvTests, spatzBenchmarks, snRuntime)",
        stim="Run the full regression on each configuration set; add the sets currently "
             "never exercised (CS-4IPU, CS-NOBUF, CS-LAT).",
        status="In Progress",
        note=notes(("spatz_pkg.sv.tpl", "// Number of IPUs in each VFU (between 1 and 8)"),
                   ("spatz_pkg.sv.tpl", "// Number of FPUs in each VFU (between 1 and 8)")),
        old="local_ci.sh runs riscvTests on CS-32B, CS-DEF, CS-SVRF, CS-DBW and "
            "spatzBenchmarks also on CS-VTL. Not covered: N_IPU != 1, N_FPU != 4, BUF_FPU off, "
            "FPU latencies (identical in every cfg).",
        cfg="all configuration sets", check="constrained-random regression"),
    dict(
        f="DOUBLE_BW VLSU",
        goal="Verify all memory items with two VLSU interfaces (2*N_FU ports).",
        test="util/local_ci.sh with spatz_cluster.doublebw.dram",
        stim="Re-run the memory items on CS-DBW; additionally stress VLSU1 write buffering.",
        status="In Progress",
        note=notes((S, "Uses `spatz_doublebw_vlsu` (two VLSU interfaces, `2*N_FU` ports). "
                       "Requires a package generated with `double_bw` / `spatz_nports = "
                       "2*n_fpu`")),
        old="riscvTests and spatzBenchmarks run on the doublebw config. Not covered: "
            "interface-1 buffering by construction.",
        cfg="CS-DBW", check="scoreboard"),
    dict(
        f="Ventaglio extension",
        goal="Verify VTL CSRs, vfxmacc.vrf/vfxmul.vrf/vventclr and redirected loads/stores.",
        test="spatzBenchmarks sp-SpMV, sp-SpMM (ventaglio cfg)",
        stim="Directed and random sparse kernels, every vtlratio/vtlidxw value, vventclr "
             "between accumulator groups.",
        status="In Progress",
        note=notes((S, "The default standalone configuration does **not** enable it.")),
        old="Only the sparse benchmarks on spatz_cluster.ventaglio.dram. Not covered: CSR "
            "semantics, every ratio/index width.",
        cfg="CS-VTL", check="scoreboard"),
]

# ----------------------------------------------------------------------------
# Configuration sheet
# ----------------------------------------------------------------------------
cfg_params = [
    # Name, Type, Definition, Legal range, Default, Knob, Old regression, Planned, Status
    ("VLEN", "localparam int unsigned (spatz_pkg)", "Bits per vector register",
     "power of two, >= N_FU*ELEN", "512", "vlen", "512, 256", "256, 512, 1024", "Swept"),
    ("N_FPU", "localparam int unsigned (spatz_pkg)", "FPUs in the VFU",
     "0..8, N_FU power of two", "4", "n_fpu", "4", "2, 4, 8", "Fixed - untested"),
    ("N_IPU", "localparam int unsigned (spatz_pkg)", "Integer units in the VFU",
     "1..8", "1", "n_ipu", "1", "1, 4", "Fixed - untested"),
    ("RVF / RVD (ELEN)", "localparam bit (spatz_pkg)", "FP32/FP64 support, ELEN = RVD ? 64 : 32",
     "isa with f / d", "1 / 1 (ELEN 64)", "cores[0].isa, spatz_fpu", "rv32imafd, rv32imaf",
     "rv32imafd, rv32imaf", "Swept"),
    ("NumVLSUInterfaces / double_bw", "localparam + define DOUBLE_BW",
     "VLSU interfaces (spatz_nports / n_fpu)", "1, 2", "1", "double_bw, spatz_nports", "1, 2",
     "1, 2", "Swept"),
    ("NrParallelInstructions", "localparam (spatz_pkg)", "Max in-flight vector instructions",
     "4 (8 with VENTAGLIO)", "4", "ventaglio", "4, 8", "4, 8", "By construction"),
    ("NrVRFBanks", "localparam (spatz_pkg)", "VRF banks", "4 (hard-coded)", "4", "-", "4", "4",
     "By construction"),
    ("mempool", "template switch", "MemPool integration (out-of-order memory, extra ports)",
     "0, 1", "0", "mempool", "0", "0 (out of scope)", "Fixed - untested"),
    ("NrMemPorts", "parameter int unsigned (spatz)", "VLSU memory ports", "equal to N_FU, power of two",
     "4", "spatz_nports", "4, 8", "equal to N_FU", "By construction"),
    ("RegisterRsp", "parameter bit (spatz)", "Unused by the controller", "0, 1", "1",
     "timing.register_offload_rsp", "1", "0, 1", "Fixed - untested"),
    ("NumOutstandingLoads", "parameter int unsigned (spatz)", "Outstanding FP LSU requests",
     ">= 1", "4", "cores[i].num_spatz_outstanding_loads", "4", "1, 4, 8", "Fixed - untested"),
    ("FPUImplementation.PipeRegs", "fpnew_pkg::fpu_implementation_t",
     "FPnew pipeline registers per opgroup and format", ">= 0",
     "ADDMUL 1/2/0/0/0/0, NONCOMP 1, CONV 2, DOTP 2", "timing.lat_*", "default only",
     "0, default, 4", "Fixed - untested"),
    ("FPUImplementation.PipeConfig", "fpnew_pkg::pipe_config_t", "Pipeline register placement",
     "BEFORE, AFTER, INSIDE, DISTRIBUTED", "BEFORE", "timing.fpu_pipe_config", "BEFORE",
     "BEFORE, DISTRIBUTED", "Fixed - untested"),
    ("FPUImplementation.UnitTypes", "fpnew_pkg::fpu_implementation_t", "FPnew unit types",
     "hard-coded in wrapper", "MERGED/DISABLED/PARALLEL/MERGED/MERGED", "-", "default only",
     "default", "By construction"),
    ("BUF_FPU", "define", "FIFO on the VFU -> VRF write path", "defined / not defined",
     "defined (buf_fpu: 1)", "buf_fpu", "defined", "defined, not defined", "Fixed - untested"),
    ("VENTAGLIO", "define", "Ventaglio sparse unit", "defined / not defined", "not defined",
     "ventaglio", "both (CS-VTL benchmarks only)", "both", "Swept"),
    ("data_width", "type width", "Width of data_t (must equal ELEN)", "equal to ELEN", "64", "data_width",
     "64, 32", "equal to ELEN", "By construction"),
    ("addr_width", "type width", "Width of addr_t", "32 recommended", "32", "addr_width", "32",
     "32", "Fixed - untested"),
]
cfg_sets = [
    ("CS-DEF", "spatz_cluster.default.dram.hjson", "Default configuration", "yes"),
    ("CS-SVRF", "spatz_cluster.smallvrf.dram.hjson", "VLEN = 256", "yes"),
    ("CS-32B", "spatz_cluster.32b.dram.hjson", "rv32imaf, ELEN = 32, data_width = 32", "yes"),
    ("CS-DBW", "spatz_cluster.doublebw.dram.hjson", "double_bw, spatz_nports = 8", "yes"),
    ("CS-VTL", "spatz_cluster.ventaglio.dram.hjson", "VENTAGLIO", "benchmarks only"),
    ("CS-4IPU", "spatz_cluster.4ipu.dram.hjson", "n_ipu = 4", "no (exists, not in local_ci.sh)"),
    ("CS-NOBUF", "to be created", "default with buf_fpu = 0", "no"),
    ("CS-LAT", "to be created", "default with different timing.lat_* and fpu_pipe_config", "no"),
]

# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {"Not Started": "F8CBAD", "In Progress": "FFE699", "Passed": "C6E0B4"}


def text_cell(ws, row, col, value):
    """Write a text cell. openpyxl turns any string starting with '=' into a formula,
    so plain text must never start with it."""
    if isinstance(value, str) and value.startswith("="):
        sys.exit(f"[vplan] text starting with '=' would become a formula: {value!r}")
    return ws.cell(row=row, column=col, value=value)

wb = Workbook()
ws = wb.active
ws.title = "Verification Plan"
cols = [
    ("Item ID", 9, "Identifier of the plan item. Stable and never reused."),
    ("Feature ID", 10, "Identifier of the feature. One feature can be covered by more than one item."),
    ("Requirement", 26, "The feature, named in a few words. [assumed] = reconstructed from the RTL, "
                        "see Known limitations in SPATZ_VERIFICATION_PLAN.md."),
    ("Verification Goal", 46, "One sentence, in the form 'Verify that ...'."),
    ("Test Name", 30, "Test, property or regression covering the item; '-' while it does not exist."),
    ("Stimulus", 52, "The scenario in operational terms: ports, traffic, order, backpressure."),
    ("Status", 13, "Not Started / In Progress / Passed. Changed only by the verification "
                   "engineer, Passed only on evidence."),
    ("Note", 60, "Verbatim quote of the specification text the requirement comes from."),
    ("Old testbench situation", 52, "What the existing environment covers, by which mechanism, "
                                    "and what it does not."),
    ("Configurations", 20, "Configuration sets of the Configuration sheet; empty = all legal "
                           "configurations."),
    ("Check", 24, "scoreboard, assertion, VIP + formal, directed test, constrained-random "
                  "regression, review, or a combination."),
]
for c, (name, width, help_) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, WRAP, BORDER
    cell.comment = Comment(help_, "vplan")
    ws.column_dimensions[cell.column_letter].width = width

for i, it in enumerate(items):
    row = [f"T{i:03d}", f"F{i:03d}", it["f"], it["goal"], it["test"], it["stim"], it["status"],
           it["note"], it["old"], it["cfg"], it["check"]]
    for c, v in enumerate(row, 1):
        cell = text_cell(ws, i + 2, c, v)
        cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
    st = ws.cell(row=i + 2, column=7)
    st.fill = PatternFill("solid", fgColor=STATUS_FILL[it["status"]])

last = len(items) + 1
dv = DataValidation(type="list", formula1='"Not Started,In Progress,Passed"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"G2:G{last}")
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:K{last}"

# Status summary (formulas)
r0 = last + 2
ws.cell(row=r0, column=3, value="Status summary").font = Font(name=FONT, bold=True, size=10)
for k, s in enumerate(["Not Started", "In Progress", "Passed"]):
    ws.cell(row=r0 + 1 + k, column=3, value=s).font = BODY
    ws.cell(row=r0 + 1 + k, column=4, value=f'=COUNTIF($G$2:$G${last},C{r0 + 1 + k})').font = BODY
ws.cell(row=r0 + 4, column=3, value="Total").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=r0 + 4, column=4, value=f"=COUNTA($A$2:$A${last})").font = Font(name=FONT, bold=True, size=10)

# Configuration sheet
wc = wb.create_sheet("Configuration")
ccols = [("Name", 28), ("Type", 30), ("Definition", 40), ("Legal range", 26),
         ("Default (spatz_cluster.default.dram.hjson)", 30), ("Knob (hjson key)", 28),
         ("Old regression values", 24), ("Planned sweep", 22), ("Verification status", 20)]
for c, (name, width) in enumerate(ccols, 1):
    cell = wc.cell(row=1, column=c, value=name)
    cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, WRAP, BORDER
    wc.column_dimensions[cell.column_letter].width = width
wc.cell(row=1, column=9).comment = Comment(
    "Swept = more than one value in the regression. Fixed - untested = a single value in every "
    "run. By construction = illegal values cannot be produced (derived or hard-coded).", "vplan")
SFILL = {"Swept": "C6E0B4", "Fixed - untested": "F8CBAD", "By construction": "DDEBF7"}
for i, p in enumerate(cfg_params):
    for c, v in enumerate(p, 1):
        cell = text_cell(wc, i + 2, c, v)
        cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
    wc.cell(row=i + 2, column=9).fill = PatternFill("solid", fgColor=SFILL[p[8]])

r = len(cfg_params) + 3
wc.cell(row=r, column=1, value="Configuration sets (referenced by the Configurations column)").font = \
    Font(name=FONT, bold=True, size=10)
for c, name in enumerate(["Set ID", "Source (hw/system/spatz_cluster/cfg/)", "Difference from CS-DEF",
                          "In old regression (util/local_ci.sh)"], 1):
    cell = wc.cell(row=r + 1, column=c, value=name)
    cell.font, cell.fill, cell.alignment, cell.border = HDR_FONT, HDR_FILL, WRAP, BORDER
for i, s in enumerate(cfg_sets):
    for c, v in enumerate(s, 1):
        cell = text_cell(wc, r + 2 + i, c, v)
        cell.font, cell.alignment, cell.border = BODY, WRAP, BORDER
wc.freeze_panes = "B2"

# Formulas carry no cached value: ask Excel/LibreOffice to recompute on open.
wb.calculation.fullCalcOnLoad = True
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"[vplan] {len(items)} items, {len(cfg_params)} parameters -> {OUT}")
